import base64
import importlib.util
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.dont_write_bytecode = True
spec = importlib.util.spec_from_file_location('file_import', Path(__file__).parents[1] / 'skills/elx-magic-study/scripts/file_import.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')


class FileImportTests(unittest.TestCase):
    def setUp(self):
        # Retain synthetic artifacts: no cleanup of user files or test directories.
        self.root = Path(tempfile.mkdtemp(prefix='elx-file-import-'))

    def convert(self, name, content):
        source = self.root / name
        source.write_bytes(content)
        output = self.root / (name + '-bundle')
        result = module.prepare(source, output)
        self.assertEqual((output / result['original']['file']).read_bytes(), content)
        self.assertFalse(result['remoteSaved'])
        return result, (output / 'content.md').read_text(encoding='utf-8')

    def test_text_and_no_overwrite(self):
        result, text = self.convert('note.md', '# 合成资料\n\n原备注'.encode())
        self.assertIn('原备注', text)
        self.assertEqual(result['coverage'], 'unreviewed')
        with self.assertRaises(FileExistsError):
            module.prepare(self.root / 'note.md', self.root / 'note.md-bundle')

    def test_image_not_fake_ocr_and_empty_text(self):
        result, text = self.convert('image.png', PNG)
        self.assertEqual(text, '')
        self.assertEqual(result['coverage'], 'needs-reading')
        self.assertFalse(result['converter']['cloudOCR'])
        result, _ = self.convert('empty.txt', b'')
        self.assertEqual(result['coverage'], 'needs-reading')

    def test_audio_and_remote_input_refused(self):
        source = self.root / 'audio.mp3'
        source.write_bytes(b'not audio')
        with self.assertRaises(ValueError):
            module.prepare(source, self.root / 'audio')
        with self.assertRaises(ValueError):
            module.prepare('https://example.com/report.pdf', self.root / 'remote')
        source = self.root / 'local.txt'
        source.write_text('sample')
        with self.assertRaises(ValueError):
            module.prepare(source, self.root / 'credential', 'https://name:password@example.com/')

    def test_html_and_csv(self):
        result, text = self.convert('sample.html', b'<h1>Topic</h1><table><tr><th>Name</th></tr><tr><td>Item</td></tr></table>')
        self.assertEqual(result['converter']['name'], 'markitdown')
        self.assertIn('Topic', text)
        self.assertIn('Item', text)
        _, text = self.convert('sample.csv', b'Name,Value\nItem,12\n')
        self.assertIn('12', text)

    def test_word_and_safe_image_names(self):
        source = self.root / 'sample.docx'
        with zipfile.ZipFile(source, 'w') as z:
            z.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
            z.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
            z.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Document sample</w:t></w:r></w:p></w:body></w:document>')
            z.writestr('word/media/../../escape.png', PNG)
        result = module.prepare(source, self.root / 'word')
        self.assertIn('Document sample', (self.root / 'word/content.md').read_text())
        self.assertEqual(result['assets'][0]['file'], 'asset-001.png')
        self.assertFalse((self.root / 'escape.png').exists())
        self.assertFalse(result['assets'][0]['selected'])

    def test_powerpoint_and_excel(self):
        from pptx import Presentation
        from openpyxl import Workbook
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[1])
        slide.shapes.title.text = 'Slide sample'
        ppt.save(self.root / 'slides.pptx')
        module.prepare(self.root / 'slides.pptx', self.root / 'slides')
        self.assertIn('Slide sample', (self.root / 'slides/content.md').read_text())
        book = Workbook()
        book.active.append(['Name', 'Value'])
        book.active.append(['Item', 17])
        book.save(self.root / 'book.xlsx')
        module.prepare(self.root / 'book.xlsx', self.root / 'book')
        self.assertIn('17', (self.root / 'book/content.md').read_text())

    def test_pdf_text_and_blank_page(self):
        for name, stream in [('text', b'BT /F1 12 Tf 20 150 Td (PDF sample) Tj ET'), ('blank', b'')]:
            objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>', b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>', b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>', b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream']
            data = b'%PDF-1.4\n'
            offsets = []
            for i, obj in enumerate(objects, 1):
                offsets.append(len(data))
                data += str(i).encode()+b' 0 obj\n'+obj+b'\nendobj\n'
            start = len(data)
            data += b'xref\n0 6\n0000000000 65535 f \n'+b''.join(f'{x:010d} 00000 n \n'.encode() for x in offsets)+b'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n'+str(start).encode()+b'\n%%EOF'
            result, text = self.convert(name+'.pdf', data)
            self.assertEqual(result['coverage'], 'needs-reading' if name == 'blank' else 'unreviewed')
            if name == 'text':
                self.assertIn('PDF sample', text)


if __name__ == '__main__':
    unittest.main()
